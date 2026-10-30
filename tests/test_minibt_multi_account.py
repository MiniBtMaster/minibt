"""多账户（方案 A：策略绑账户）+ 期货账户落盘（方案 C）。

* ``Bt.addstrategy(MyStrategy, account='H宏源期货')`` → 策略 ``_live_account``
  → ``KLine(account=)`` → ``TqObjs(account=)`` → ``get_position`` /
  ``TargetPosTask`` 全部落到该账户 ✓
* ``bt.bindAccounts()`` 把名字解析成 tqsdk 账户实例（``run()`` 自动调）
* ``LiveTools`` 优先用策略绑定的账户（其次快期模拟）
* ``TqAccountStore``：期货权益曲线 + 成交落盘（sqlite），重启不丢、可多进程汇总
"""
import os
import pathlib
import subprocess
import sys
import time

import pytest

REPO = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))

from minibt import Bt, Strategy                          # noqa: E402
from minibt.strategy.live_tools import (                 # noqa: E402
    LiveTools, _account_label)


# ------------------------------------------------------------------ 假账户
class FakeAccount:
    def __init__(self, broker_id='', account_id=''):
        self.broker_id = broker_id
        self.account_id = account_id


class FakeApi:
    """够用的假 TqApi（多账户 / 单账户两种形态 ✓）"""

    def __init__(self, accounts=None, trades=None):
        multi = type('Multi', (), {'_account_list': list(accounts or [])})()
        self._account = multi if accounts else FakeAccount('快期', '模拟')
        self._trades = list(trades or [])
        self.account_calls = []

    def get_account(self, account=None):
        self.account_calls.append(account)
        return type('Acc', (), {'balance': 100000.0, 'available': 90000.0,
                                'float_profit': 120.0, 'close_profit': 30.0,
                                'margin': 10000.0, 'risk_ratio': 0.1})()

    def get_trade(self, account=None):
        return self._trades

    def get_position(self, symbol=None, account=None):
        return {}

    def get_order(self, account=None):
        return {}


class FakeOwner:
    _realtime = True

    def __init__(self, strategies=(), api=None):
        self.strategies = list(strategies)
        self.chart = object()
        self.window = type('Win', (), {'contract': 'SHFE.rb2610_60'})()
        self._api = api

    def _first_api(self):
        return self._api

    def _strategy_obj(self):
        return self.strategies[0] if self.strategies else None


class MyStrategy(Strategy):
    def next(self):
        pass


# ======================================================== 方案 A：账户绑定
def _bt():
    """不跑 ``Bt.__init__``（会去扫数据），只测 addstrategy / bindAccounts ✓"""
    bt = Bt.__new__(Bt)
    bt.strategies = []
    bt._api = None
    return bt


def test_addstrategy_binds_account():
    bt = _bt()
    bt.addstrategy(MyStrategy, account='H宏源期货')
    assert getattr(MyStrategy, '_live_account') == 'H宏源期货'
    # account= 没被当成策略参数（Strategy.account 是基类的 property ✓）
    assert 'account' not in vars(MyStrategy)
    assert MyStrategy in bt.strategies
    # 清理类属性（避免污染其它用例）
    del MyStrategy._live_account


def test_addstrategy_keeps_params_and_supports_path(tmp_path):
    bt = _bt()
    script = tmp_path / 'demo.py'
    script.write_text(
        'from minibt import Strategy\n'
        'class FromFile(Strategy):\n'
        '    def next(self): pass\n', encoding='utf-8')
    bt.addstrategy(str(script), account='Y银河期货', period=7)
    from minibt import Strategy as Base
    found = [s for s in bt.strategies if s.__name__ == 'FromFile']
    assert found, '没加载到文件里的策略'
    assert getattr(found[0], '_live_account') == 'Y银河期货'
    del found[0]._live_account
    assert Base is Strategy


def test_bind_accounts_resolves_names():
    sim = FakeAccount('快期', '模拟')
    real = FakeAccount('H宏源期货', '123456')
    api = FakeApi([sim, real])
    bt = _bt()
    bt._api = api
    bt.addstrategy(MyStrategy, account='H宏源期货')
    bt.bindAccounts()
    assert getattr(MyStrategy, '_live_account') is real      # 名字 → 实例 ✓
    del MyStrategy._live_account

    # account_id 也认，而且别名 bindaccounts 一致 ✓
    bt2 = _bt()
    bt2._api = api
    bt2.addstrategy(MyStrategy, account='123456')
    bt2.bindaccounts()
    assert getattr(MyStrategy, '_live_account') is real
    del MyStrategy._live_account

    # 找不到 → 抛错（把可选账户打出来）
    bt3 = _bt()
    bt3._api = api
    bt3.addstrategy(MyStrategy, account='不存在')
    with pytest.raises(ValueError):
        bt3.bindAccounts()
    del MyStrategy._live_account


def test_bind_accounts_single_account_accepts_any_name():
    api = FakeApi()                     # 单账户（没有 _account_list）
    bt = _bt()
    bt._api = api
    bt.addstrategy(MyStrategy, account='随便写')
    bt.bindAccounts()
    assert getattr(MyStrategy, '_live_account') is api._account
    del MyStrategy._live_account


def test_live_tools_prefers_strategy_account():
    sim = FakeAccount('快期', '模拟')
    real = FakeAccount('H宏源期货', '123456')
    api = FakeApi([sim, real])
    owner = FakeOwner([MyStrategy], api)
    MyStrategy._live_account = real
    try:
        tools = LiveTools(owner)
        assert tools._trade_account() is real            # 策略绑的优先 ✓
        assert tools._resolve_account('H宏源期货') is real
        assert tools._resolve_account('missing') is None  # 多账户找不到 → None
        assert tools._resolve_account(real) is real       # 实例原样返回 ✓
    finally:
        del MyStrategy._live_account


def test_account_label():
    assert _account_label(FakeAccount('H宏源期货', '123456')) == 'H宏源期货/123456'
    assert _account_label(FakeAccount('H宏源期货', '')) == 'H宏源期货'
    assert _account_label(None) == ''
    assert _account_label(type('TqKq', (), {})()) == '快期模拟'


# ==================================================== 方案 C：期货账户落盘
class FakeTrade:
    def __init__(self, trade_id, price=3600.0, volume=1):
        self.trade_id = trade_id
        self.trade_date_time = '2026-10-18 09:31:00'
        self.instrument_id = 'SHFE.rb2610'
        self.direction = 'BUY'
        self.offset = 'OPEN'
        self.price = price
        self.volume = volume
        self.commission = 1.5
        self.realized_pnl = 0.0


def _tools_with_store(tmp_path, *, kind='期货', trades=()):
    from minibt.tq_store import TqAccountStore

    api = FakeApi([FakeAccount('H宏源期货', '123456')], list(trades))
    owner = FakeOwner([], api)
    tools = LiveTools(owner)
    tools._store_obj = TqAccountStore(str(tmp_path / 'tq.sqlite3'))
    tools._store_at = -1e9                       # 不节流 ✓
    summary = {'kind': kind, 'account': 'H宏源期货/123456',
               'balance': 100000.0, 'available': 90000.0,
               'float_profit': 120.0, 'close_profit': 30.0,
               'extra': 10000.0, 'ratio': 10.0}
    return tools, api, summary


def test_store_tick_persists_equity_and_trades(tmp_path):
    tools, _api, summary = _tools_with_store(
        tmp_path, trades=[FakeTrade('T1'), FakeTrade('T2')])
    tools._store_tick(summary)
    store = tools.account_store()
    assert store is not None
    equity = store.equity('H宏源期货/123456')
    assert len(equity) == 1 and equity[0][1] == 100000.0
    assert [t['trade_id'] for t in store.trades('H宏源期货/123456')] == ['T2', 'T1']

    # 再落一次：权益点新增（时间不同）、成交不重复 ✓
    time.sleep(1.05)                           # 权益点按**秒**去重
    tools._store_at = -1e9
    tools._store_tick(dict(summary, balance=101000.0))
    assert len(store.equity('H宏源期货/123456')) == 2
    assert len(store.trades('H宏源期货/123456')) == 2


def test_store_tick_throttles_and_skips_stock(tmp_path):
    tools, _api, summary = _tools_with_store(tmp_path)
    tools._store_tick(summary)
    tools._store_tick(dict(summary, balance=1.0))     # 5 秒内 → 跳过 ✓
    store = tools.account_store()
    assert len(store.equity('H宏源期货/123456')) == 1

    stock_tools, _api2, stock = _tools_with_store(tmp_path / 'b', kind='股票')
    stock_tools._store_tick(stock)
    assert stock_tools.account_store().equity('H宏源期货/123456') == []


def test_store_tick_via_summary_cache(tmp_path):
    tools, _api, summary = _tools_with_store(tmp_path)
    tools._account_summary = lambda: summary          # 不碰假 api 的细节
    cached = tools._account_summary_cached()
    assert cached == summary
    assert len(tools.account_store().equity('H宏源期货/123456')) == 1


def test_store_disabled(tmp_path):
    tools, _api, summary = _tools_with_store(tmp_path)
    tools.store_enabled = False
    tools._store_tick(summary)
    assert tools.account_store() is None


def test_tutorial_overview_script(tmp_path):
    from minibt.tq_store import TqAccountStore

    db = str(tmp_path / 'tq.sqlite3')
    store = TqAccountStore(db)
    store.record_account('H宏源期货/123456', {'balance': 100000.0},
                         ts='2026-10-18 09:30:00')
    store.record_account('H宏源期货/123456', {'balance': 101500.0},
                         ts='2026-10-18 09:35:00')
    store.close()

    script = REPO / 'tutorials' / 'realtime' / '10_accounts_overview.py'
    assert script.exists()
    result = subprocess.run(
        [sys.executable, str(script), '--db', db, '--check'],
        capture_output=True, text=True, encoding='utf-8',
        env={**os.environ, 'PYTHONIOENCODING': 'utf-8'})
    assert result.returncode == 0, result.stderr[-800:]
    assert 'H宏源期货/123456' in result.stdout
    assert '101,500.00' in result.stdout


# =============================== 方案 B：多账户对比 + 账户页读落盘历史
class PerAccountApi:
    """每个账户返回不同的 balance（多账户对比用 ✓）"""

    def __init__(self, mapping):
        self.mapping = dict(mapping)
        items = [FakeAccount(broker, '1001') for broker in self.mapping]
        self._account = type('Multi', (), {'_account_list': items})()

    def get_account(self, account=None):
        broker = getattr(account, 'broker_id', '')
        balance, margin = self.mapping.get(broker, (0.0, 0.0))
        return type('Acc', (), {'balance': balance, 'available': balance * 0.9,
                                'float_profit': 10.0, 'close_profit': 0.0,
                                'margin': margin, 'risk_ratio': 0.1})()

    def get_trade(self, account=None):
        return []

    def get_position(self, symbol=None, account=None):
        return {}

    def get_order(self, account=None):
        return {}


def test_account_instances_single_and_multi():
    multi = LiveTools(FakeOwner([], PerAccountApi({'A': (1, 0), 'B': (2, 0)})))
    assert len(multi._account_instances()) == 2

    single = LiveTools(FakeOwner([], FakeApi()))
    assert len(single._account_instances()) == 1
    empty = LiveTools(FakeOwner([], None))
    assert empty._account_instances() == []


def test_account_summaries_multi_with_store_history(tmp_path):
    from minibt.tq_store import TqAccountStore

    api = PerAccountApi({'H宏源期货': (101500.0, 10500.0),
                         'Y银河期货': (49500.0, 5000.0)})
    tools = LiveTools(FakeOwner([], api))
    tools._store_obj = TqAccountStore(str(tmp_path / 'tq.sqlite3'))
    store = tools.store_obj = tools._store_obj
    for name, first, last in (('H宏源期货/1001', 100000.0, 101500.0),
                              ('Y银河期货/1001', 50000.0, 49500.0)):
        store.record_account(name, {'balance': first}, ts='2026-10-18 09:30:00')
        store.record_account(name, {'balance': last}, ts='2026-10-18 09:35:00')

    rows = tools._account_summaries()
    assert [row['account'] for row in rows] == ['H宏源期货/1001',
                                                'Y银河期货/1001']   # 权益降序 ✓
    assert rows[0]['range_profit'] == 1500.0
    assert rows[1]['range_profit'] == -500.0
    assert rows[0]['points'] == 2


def test_account_summaries_single_is_empty():
    tools = LiveTools(FakeOwner([], FakeApi()))     # 单账户 → 不画对比区
    assert tools._account_summaries() == []


def test_equity_curve_prefers_stored_history(tmp_path):
    tools, _api, summary = _tools_with_store(tmp_path)
    tools._store_obj = tools.account_store()
    store = tools._store_obj
    store.record_account('H宏源期货/123456', {'balance': 100000.0},
                         ts='2026-10-18 09:30:00')
    store.record_account('H宏源期货/123456', {'balance': 101500.0},
                         ts='2026-10-18 09:35:00')
    tools._equity_samples = [('10:00:00', 1.0)]     # 内存采样应被忽略 ✓
    curve = tools._equity_curve()
    assert curve == [('2026-10-18 09:30:00', 100000.0),
                     ('2026-10-18 09:35:00', 101500.0)]
    # 库空 → 退回内存采样 ✓
    store.reset()
    assert tools._equity_curve() == [('10:00:00', 1.0)]


def test_sync_account_page_passes_accounts(tmp_path):
    api = PerAccountApi({'H宏源期货': (101500.0, 0.0), 'Y银河期货': (49500.0, 0.0)})

    class Window:
        def __init__(self):
            self.calls = []

        def is_account_page(self):
            return True

        def update_account_page(self, **kwargs):
            self.calls.append(kwargs)

    owner = FakeOwner([], api)
    owner.window = Window()
    tools = LiveTools(owner)
    tools._store_obj = False                     # 不落盘，只测传参 ✓
    tools._sync_account_page()
    assert owner.window.calls
    accounts = owner.window.calls[-1].get('accounts')
    assert accounts and len(accounts) == 2
    assert owner.window.calls[-1].get('summary')


# ---------------------------------- AccountPage：多账户对比表（纯渲染）
class _Panel:
    id = 'window.__p'

    def __init__(self):
        self.html = ''

    def set_html(self, html):
        self.html = html

    def run_script(self, script):
        pass

    def show(self):
        pass

    def hide(self):
        pass


class _Chart:
    def __init__(self):
        self.panel = _Panel()
        self.win = type('W', (), {'handlers': {}})()

    def create_html_panel(self):
        return self.panel


def _row(name, balance, profit=None):
    row = {'kind': '期货', 'account': name, 'balance': balance,
           'available': balance * 0.9, 'float_profit': 10.0,
           'close_profit': 0.0, 'extra_label': '保证金', 'extra': 1000.0,
           'ratio_label': '风险度', 'ratio': 10.0, 'ratio_percent': True,
           'positions': 1, 'orders': 0}
    if profit is not None:
        row['range_profit'] = profit
    return row


def test_account_page_multi_account_section(monkeypatch, tmp_path):
    from minibt.strategy import pages as pages_mod

    monkeypatch.setattr(pages_mod, 'SORT_FILE', str(tmp_path / 'p.json'))
    chart = _Chart()
    page = pages_mod.AccountPage(chart)

    page.update(summary=_row('H宏源期货/1001', 101500.0))
    assert '多账户对比' not in chart.panel.html       # 单账户 → 不显示 ✓

    page.update(accounts=[_row('H宏源期货/1001', 101500.0, 1500.0),
                          _row('Y银河期货/1001', 49500.0, -500.0)])
    html = chart.panel.html
    assert '多账户对比' in html and '2 个账户' in html
    assert 'H宏源期货/1001' in html and 'Y银河期货/1001' in html
    assert '1,500.00' in html and '-500.00' in html
    assert '151,000.00' in html                      # 合计 ✓

    page.update(accounts=[_row('H宏源期货/1001', 101500.0)])
    assert '多账户对比' not in chart.panel.html       # 又只有 1 个 → 收起 ✓


# ---------------------- 多账户对比表：排序 / 导出 / 保值 + 多账户曲线叠加
def test_compare_table_sort_and_export(monkeypatch, tmp_path):
    import json

    from minibt.strategy import pages as pages_mod

    monkeypatch.setattr(pages_mod, 'SORT_FILE', str(tmp_path / 'p.json'))
    monkeypatch.setattr(pages_mod, 'EXPORTS_DIR', str(tmp_path))
    chart = _Chart()
    page = pages_mod.AccountPage(chart)
    rows = [_row('H宏源期货/1001', 101500.0, 1500.0),
            _row('Y银河期货/1001', 49500.0, -500.0)]
    page.update(summary=rows[0], accounts=rows)

    page.sort_by('accounts', 1)                       # 按权益升序 ✓
    assert page.tables['accounts']['rows'][0][0] == 'Y银河期货/1001'
    page.sort_by('accounts', 1)                       # 再点 → 降序 ✓
    assert page.tables['accounts']['rows'][0][0] == 'H宏源期货/1001'
    saved = json.loads(pathlib.Path(pages_mod.SORT_FILE).read_text('utf-8'))
    assert saved['sort']['accounts'] == [1, False]

    page.update(accounts=rows)                        # 刷新后仍按记忆排序 ✓
    assert page.tables['accounts']['rows'][0][0] == 'H宏源期货/1001'

    path = page.export_csv('accounts')                # CSV 全部行 ✓
    assert path and pathlib.Path(path).exists()
    assert '多账户对比' in pathlib.Path(path).name
    text = pathlib.Path(path).read_text(encoding='utf-8-sig')
    assert 'H宏源期货/1001' in text and '101500.0' in text

    page.update(tables={'positions': {'columns': ('代码',), 'rows': [['rb2610']],
                                      'total': 1, 'limit': 1}})
    book = page.export_all()                          # xlsx 里也有这一页 ✓
    assert book and book.endswith('.xlsx')
    from openpyxl import load_workbook
    assert '多账户对比' in load_workbook(book).sheetnames


def test_multi_curve_overlay(monkeypatch, tmp_path):
    import time as _time

    from minibt.strategy import pages as pages_mod

    monkeypatch.setattr(pages_mod, 'SORT_FILE', str(tmp_path / 'p.json'))
    chart = _Chart()
    page = pages_mod.AccountPage(chart)
    now = _time.time()
    fmt = lambda minutes: _time.strftime(  # noqa: E731
        '%Y-%m-%d %H:%M:%S', _time.localtime(now - minutes * 60))
    curves = {'A/1': [(fmt(9), 100000.0), (fmt(3), 101500.0)],
              'B/1': [(fmt(9), 50000.0), (fmt(3), 49500.0)]}
    page.update(account_curves=curves)
    html = chart.panel.html
    assert '多账户收益对比' in html
    assert '+1.50%' in html and '-1.00%' in html
    assert html.count('<polyline class="line"') == 2
    assert 'legend-item' in html

    page.update(account_curves={'A/1': curves['A/1']})   # 单条 → 不画 ✓
    assert '多账户收益对比' not in chart.panel.html
    page.update(account_curves={'A/1': [(fmt(5), 0.0), (fmt(1), 1.0)]})
    assert '多账户收益对比' not in chart.panel.html      # 首点为 0 → 跳过 ✓


def test_account_curves_from_store(tmp_path):
    from minibt.tq_store import TqAccountStore

    api = PerAccountApi({'H宏源期货': (101500.0, 0.0),
                         'Y银河期货': (49500.0, 0.0)})
    tools = LiveTools(FakeOwner([], api))
    tools._store_obj = TqAccountStore(str(tmp_path / 'tq.sqlite3'))
    store = tools._store_obj
    for name, first, last in (('H宏源期货/1001', 100000.0, 101500.0),
                              ('Y银河期货/1001', 50000.0, 49500.0)):
        store.record_account(name, {'balance': first}, ts='2026-10-18 09:30:00')
        store.record_account(name, {'balance': last}, ts='2026-10-18 09:35:00')
    curves = tools._account_curves()
    assert sorted(curves) == ['H宏源期货/1001', 'Y银河期货/1001']
    assert curves['H宏源期货/1001'][-1][1] == 101500.0

    single = LiveTools(FakeOwner([], FakeApi(accounts=[FakeAccount('A', '1')])))
    assert single._account_curves() == {}             # 单账户 → 不叠 ✓
